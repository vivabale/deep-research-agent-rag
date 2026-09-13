from pathlib import Path
from docx import Document as DocxDocument
from ..schemas import Document
from .base import BaseLoader

class DocxLoader(BaseLoader):
    def load(self,path):
        p=Path(path)
        doc=DocxDocument(path)
        return Document("\n".join(x.text for x in doc.paragraphs), {"source":p.name,"filename":p.name,"filetype":p.suffix.lower()})
